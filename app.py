import datetime
import base64
import os
import pathlib
import shutil
import tempfile
import re
import zipfile
import xml.etree.ElementTree as ET
from io import BytesIO
from pathlib import Path
import pandas as pd
import streamlit as st

# محاولة استيراد مكتبة الرسوم البيانية بأمان
try:
  import plotly.express as px

  HAS_PLOTLY = True
except ImportError:
  HAS_PLOTLY = False

# =====================================================
# 1. إعدادات التطبيق وتسجيل الدخول
# =====================================================
st.set_page_config(
    page_title="OneSpot Cafe - Financial System", page_icon="☕", layout="wide"
)

APP_PASSWORD = "Cofe2026"
CURRENCY_LABEL = "ج.م"
BRANCHES = ["Onespot 1", "OneSpot 2"]
LEGACY_BRANCH = BRANCHES[0]


def format_currency(amount):
  return f"{amount:,.2f} {CURRENCY_LABEL}"

DATA_FILE = Path("cafe_data.csv")
BACKUP_FILE = Path("cafe_data.csv.bak")

COLUMNS = ["ID", "Date", "Month_No", "Branch", "Category", "Type", "Amount", "Description"]

INCOME_CATS = [
    "Daily Sales (المبيعات اليوميه)",
    "Other Income (إيرادات أخرى)",
]

EXPENSE_CATS = [
    "Coffee Beans & Milk (مواد خام)",
    "Staff Salaries (رواتب الموظفين)",
    "Rent & Utilities (إيجار وفواتير)",
    "(valu installments) (اقساط فاليو)",
    "Packaging (أكواب ومطبوعات)",
    "Other Expense (مصروفات أخرى)",
]


def password_entered():
    entered = st.session_state.get("password", "")
    st.session_state["password_correct"] = entered == APP_PASSWORD


def apply_login_background():
  """Show a local cafe photo on the login screen, with a safe CSS fallback."""
  image_path = Path(__file__).resolve().parent / "cafe_login_bg.png"
  try:
    image_data = base64.b64encode(image_path.read_bytes()).decode("ascii")
    background = (
        "linear-gradient(rgba(8, 10, 12, 0.68), rgba(8, 10, 12, 0.72)), "
        f"url('data:image/png;base64,{image_data}')"
    )
  except (OSError, ValueError):
    background = "linear-gradient(135deg, #17120e, #3b2a1d)"

  st.markdown(
      f"""
      <style>
      .stApp, [data-testid="stAppViewContainer"] {{
          background-image: {background};
          background-size: contain;
          background-position: center;
          background-repeat: no-repeat;
          background-color: #17120e;
          background-attachment: fixed;
      }}
      [data-testid="stHeader"] {{ background: rgba(0, 0, 0, 0); }}
      [data-testid="stMainBlockContainer"] {{
          background: rgba(8, 10, 12, 0.16);
          border-radius: 16px;
      }}
      [data-testid="stMainBlockContainer"] h1,
      [data-testid="stMainBlockContainer"] label {{
          color: #ffffff !important;
          text-shadow: 0 2px 8px rgba(0, 0, 0, 0.9);
      }}
      [data-testid="stMainBlockContainer"] input {{
          background-color: rgba(255, 255, 255, 0.96);
          color: #171717;
      }}
      </style>
      """,
      unsafe_allow_html=True,
  )


def check_password():
    if st.session_state.get("password_correct", False):
        return True

    apply_login_background()
    login_banner = Path(__file__).resolve().parent / "login_banner.png"
    if login_banner.is_file():
        _, banner_col, _ = st.columns([1, 2.4, 1])
        with banner_col:
            st.image(str(login_banner), use_container_width=True)
    if not login_banner.is_file():
        st.title("🔒 OneSpot Cafe System - Secure Login")

    st.text_input(
        "Enter Cafe Password:",
        type="password",
        key="password",
        on_change=password_entered,
    )

    if st.session_state.get("password_correct") is False:
        st.error("❌ Incorrect password. Please try again.")

    return False


if not check_password():
  st.stop()


# =====================================================
# 2. قراءة البيانات والتحقق منها
# =====================================================
def load_csv_data():
  if not DATA_FILE.exists():
    return pd.DataFrame(columns=COLUMNS), None

  try:
    df = pd.read_csv(DATA_FILE)
  except Exception:
    return None, (
        "تعذر قراءة ملف cafe_data.csv. تم إيقاف التعديلات لحماية بياناتك."
    )

  if df.empty and len(df.columns) == 0:
    return pd.DataFrame(columns=COLUMNS), None

  if "ID" not in df.columns:
    df.insert(0, "ID", range(1, len(df) + 1))

  for col in COLUMNS:
    if col not in df.columns:
      df[col] = LEGACY_BRANCH if col == "Branch" else pd.NA

  df["Branch"] = df["Branch"].fillna(LEGACY_BRANCH).replace("", LEGACY_BRANCH)

  df = df[COLUMNS].copy()

  df["ID"] = pd.to_numeric(df["ID"], errors="coerce")
  df["Amount"] = pd.to_numeric(df["Amount"], errors="coerce")
  parsed_dates = pd.to_datetime(df["Date"], errors="coerce")

  valid = (
      df["ID"].notna()
      & (df["ID"] > 0)
      & (df["ID"] % 1 == 0)
      & parsed_dates.notna()
      & df["Category"].notna()
      & df["Type"].isin(["Income", "Expense"])
      & df["Amount"].notna()
      & (df["Amount"] > 0)
  )

  invalid_count = int((~valid).sum())

  if invalid_count:
    return None, (
        f"تم العثور على {invalid_count} صف غير صالح في ملف البيانات. تم إيقاف"
        " التعديلات حتى مراجعة الملف الأصلي."
    )

  df["ID"] = df["ID"].astype("int64")

  if df["ID"].duplicated().any():
    return None, (
        "يوجد تكرار في أرقام ID. تم إيقاف الحذف والحفظ لتجنب تعديل القيد الخطأ."
    )

  df["Date"] = parsed_dates.dt.strftime("%Y-%m-%d")
  df["Month_No"] = parsed_dates.dt.month
  df["Branch"] = df["Branch"].astype(str)
  df["Category"] = df["Category"].astype(str)
  df["Type"] = df["Type"].astype(str)
  df["Amount"] = df["Amount"].astype(float)
  df["Description"] = df["Description"].fillna("").astype(str)

  return df[COLUMNS].reset_index(drop=True), None


# =====================================================
# 3. حفظ البيانات والنسخة الاحتياطية
# =====================================================
def save_csv_data(df):
  df = df[COLUMNS].copy()

  df["ID"] = pd.to_numeric(df["ID"], errors="raise")
  df["Amount"] = pd.to_numeric(df["Amount"], errors="raise")
  parsed_dates = pd.to_datetime(df["Date"], errors="raise")

  if df["ID"].isna().any() or df["ID"].duplicated().any():
    raise ValueError("أرقام القيود مفقودة أو مكررة.")

  if (df["ID"] <= 0).any() or (df["ID"] % 1 != 0).any():
    raise ValueError("يوجد رقم قيد غير صالح.")

  if df["Amount"].isna().any() or (df["Amount"] <= 0).any():
    raise ValueError("يوجد مبلغ غير صالح.")

  if not df["Type"].isin(["Income", "Expense"]).all():
    raise ValueError("يوجد نوع معاملة غير صالح.")

  df["ID"] = df["ID"].astype("int64")
  df["Date"] = parsed_dates.dt.strftime("%Y-%m-%d")
  df["Month_No"] = parsed_dates.dt.month
  df["Branch"] = df["Branch"].fillna(LEGACY_BRANCH).astype(str)

  DATA_FILE.parent.mkdir(parents=True, exist_ok=True)

  if DATA_FILE.exists():
    shutil.copy2(DATA_FILE, BACKUP_FILE)

  temp_path = None

  try:
    with tempfile.NamedTemporaryFile(
        mode="w",
        encoding="utf-8-sig",
        newline="",
        suffix=".tmp",
        prefix="cafe_data_",
        dir=str(DATA_FILE.parent),
        delete=False,
    ) as temp_file:
      temp_path = Path(temp_file.name)
      df.to_csv(temp_file, index=False)
      temp_file.flush()
      os.fsync(temp_file.fileno())

    os.replace(temp_path, DATA_FILE)
    temp_path = None

  finally:
    if temp_path is not None and temp_path.exists():
      temp_path.unlink()


def read_database_url():
  try:
    direct_url = st.secrets.get("DATABASE_URL", "")
    if direct_url:
      return str(direct_url).strip()

    # Also accept the nested TOML layout some Streamlit users configure.
    connections = st.secrets.get("connections", {})
    postgresql = connections.get("postgresql", {}) if hasattr(connections, "get") else {}
    nested_url = postgresql.get("DATABASE_URL", "") or postgresql.get("url", "")
    return str(nested_url).strip()
  except Exception:
    return ""


@st.cache_resource(show_spinner=False)
def create_database_engine(database_url):
  from sqlalchemy import create_engine

  return create_engine(database_url, pool_pre_ping=True, pool_recycle=300)


def get_database_engine():
  database_url = read_database_url()
  return create_database_engine(database_url) if database_url else None


def ensure_database(engine):
  from sqlalchemy import text

  with engine.begin() as connection:
    # Serialize first-time setup so simultaneous visits cannot import the CSV twice.
    connection.execute(text("SELECT pg_advisory_xact_lock(784213905)"))
    connection.execute(text("""
      CREATE TABLE IF NOT EXISTS transactions (
        "ID" BIGINT PRIMARY KEY,
        "Date" DATE NOT NULL,
        "Month_No" INTEGER NOT NULL,
        "Branch" TEXT NOT NULL,
        "Category" TEXT NOT NULL,
        "Type" TEXT NOT NULL CHECK ("Type" IN ('Income', 'Expense')),
        "Amount" DOUBLE PRECISION NOT NULL CHECK ("Amount" > 0),
        "Description" TEXT NOT NULL DEFAULT ''
      )
    """))
    connection.execute(text("""
      CREATE TABLE IF NOT EXISTS onespot_metadata (
        key TEXT PRIMARY KEY,
        value TEXT NOT NULL
      )
    """))
    connection.execute(text("""
      CREATE TABLE IF NOT EXISTS attendance_records (
        employee TEXT NOT NULL,
        work_date DATE NOT NULL,
        day_label TEXT,
        clock_in TEXT,
        clock_out TEXT,
        hours DOUBLE PRECISION NOT NULL DEFAULT 0,
        salary_advance DOUBLE PRECISION NOT NULL DEFAULT 0,
        team TEXT NOT NULL DEFAULT 'OneSpot',
        PRIMARY KEY (employee, work_date)
      )
    """))
    imported = connection.execute(
        text("SELECT value FROM onespot_metadata WHERE key = 'csv_import_v1'")
    ).scalar_one_or_none()
    if imported is None:
      legacy_df, legacy_error = load_csv_data()
      if legacy_error:
        raise ValueError(legacy_error)
      if not legacy_df.empty:
        # PostgreSQL DATE columns must receive Python date values, not CSV strings.
        legacy_df["Date"] = pd.to_datetime(legacy_df["Date"], errors="raise").dt.date
        legacy_df.to_sql("transactions", con=connection, if_exists="append", index=False)
      connection.execute(
          text("INSERT INTO onespot_metadata (key, value) VALUES ('csv_import_v1', 'done')")
      )


def load_data():
  engine = get_database_engine()
  if engine is None:
    return load_csv_data()
  try:
    from sqlalchemy import text

    ensure_database(engine)
    df = pd.read_sql_query(
        text('SELECT "ID", "Date", "Month_No", "Branch", "Category", "Type", "Amount", "Description" FROM transactions ORDER BY "ID"'),
        engine,
    )
    if df.empty:
      return pd.DataFrame(columns=COLUMNS), None
    parsed_dates = pd.to_datetime(df["Date"], errors="coerce")
    df["ID"] = pd.to_numeric(df["ID"], errors="coerce").astype("int64")
    df["Amount"] = pd.to_numeric(df["Amount"], errors="coerce").astype(float)
    df["Date"] = parsed_dates.dt.strftime("%Y-%m-%d")
    df["Month_No"] = parsed_dates.dt.month
    df["Branch"] = df["Branch"].fillna(LEGACY_BRANCH).astype(str)
    df["Category"] = df["Category"].astype(str)
    df["Type"] = df["Type"].astype(str)
    df["Description"] = df["Description"].fillna("").astype(str)
    return df[COLUMNS].reset_index(drop=True), None
  except Exception as exc:
    return None, f"تعذر الاتصال بقاعدة Neon أو قراءة القيود منها: {exc}"


def save_data(df):
  df = df[COLUMNS].copy()
  df["ID"] = pd.to_numeric(df["ID"], errors="raise")
  df["Amount"] = pd.to_numeric(df["Amount"], errors="raise")
  parsed_dates = pd.to_datetime(df["Date"], errors="raise")
  if df["ID"].isna().any() or df["ID"].duplicated().any():
    raise ValueError("أرقام القيود مفقودة أو مكررة.")
  if (df["ID"] <= 0).any() or (df["ID"] % 1 != 0).any():
    raise ValueError("يوجد رقم قيد غير صالح.")
  if df["Amount"].isna().any() or (df["Amount"] <= 0).any():
    raise ValueError("يوجد مبلغ غير صالح.")
  if not df["Type"].isin(["Income", "Expense"]).all():
    raise ValueError("يوجد نوع معاملة غير صالح.")
  df["ID"] = df["ID"].astype("int64")
  # Bind as a PostgreSQL DATE instead of sending VARCHAR values.
  df["Date"] = parsed_dates.dt.date
  df["Month_No"] = parsed_dates.dt.month
  df["Branch"] = df["Branch"].fillna(LEGACY_BRANCH).astype(str)

  engine = get_database_engine()
  if engine is None:
    return save_csv_data(df)
  from sqlalchemy import text

  ensure_database(engine)
  with engine.begin() as connection:
    connection.execute(text("DELETE FROM transactions"))
    if not df.empty:
      df.to_sql("transactions", con=connection, if_exists="append", index=False)


def save_attendance_data(df):
  engine = get_database_engine()
  if engine is None or df.empty:
    return
  from sqlalchemy import text

  ensure_database(engine)
  rows = []
  for row in df.to_dict("records"):
    work_date = pd.to_datetime(row["التاريخ"], errors="coerce")
    if pd.isna(work_date):
      continue
    rows.append({
        "employee": str(row["الموظف"]),
        "work_date": work_date.date(),
        "day_label": None if pd.isna(row.get("اليوم")) else str(row.get("اليوم")),
        "clock_in": None if pd.isna(row.get("الحضور")) else str(row.get("الحضور")),
        "clock_out": None if pd.isna(row.get("الانصراف")) else str(row.get("الانصراف")),
        "hours": float(row.get("ساعات العمل") or 0),
        "salary_advance": float(row.get("السلفة (ج.م)") or 0),
        "team": "OneSpot",
    })
  if not rows:
    return
  with engine.begin() as connection:
    connection.execute(text("""
      INSERT INTO attendance_records
        (employee, work_date, day_label, clock_in, clock_out, hours, salary_advance, team)
      VALUES
        (:employee, :work_date, :day_label, :clock_in, :clock_out, :hours, :salary_advance, :team)
      ON CONFLICT (employee, work_date) DO UPDATE SET
        day_label = EXCLUDED.day_label,
        clock_in = EXCLUDED.clock_in,
        clock_out = EXCLUDED.clock_out,
        hours = EXCLUDED.hours,
        salary_advance = EXCLUDED.salary_advance,
        team = EXCLUDED.team
    """), rows)


def load_attendance_data():
  engine = get_database_engine()
  if engine is None:
    return pd.DataFrame()
  from sqlalchemy import text

  ensure_database(engine)
  saved = pd.read_sql_query(text("""
    SELECT employee, work_date, day_label, clock_in, clock_out, hours, salary_advance, team
    FROM attendance_records
    ORDER BY work_date DESC, employee
  """), engine)
  if saved.empty:
    return pd.DataFrame()
  return saved.rename(columns={
      "employee": "الموظف",
      "work_date": "التاريخ",
      "day_label": "اليوم",
      "clock_in": "الحضور",
      "clock_out": "الانصراف",
      "hours": "ساعات العمل",
      "salary_advance": "السلفة (ج.م)",
      "team": "الفريق",
  }).assign(**{"التاريخ": lambda frame: pd.to_datetime(frame["التاريخ"], errors="coerce")})


def read_attendance_with_builtin_excel(uploaded_file):
  """Read the workbook's cached cell values with Python's standard library."""
  ns = {
      "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "pkg": "http://schemas.openxmlformats.org/package/2006/relationships",
  }
  records = []
  ignored_sheets = []

  with zipfile.ZipFile(BytesIO(uploaded_file.getvalue())) as archive:
    names = set(archive.namelist())
    shared_strings = []
    if "xl/sharedStrings.xml" in names:
      root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
      for item in root.findall("main:si", ns):
        shared_strings.append("".join(node.text or "" for node in item.findall(".//main:t", ns)))

    workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
    relationships_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    relationships = {
        relationship.attrib["Id"]: relationship.attrib["Target"]
        for relationship in relationships_root.findall("pkg:Relationship", ns)
    }

    for sheet_info in workbook_root.findall("main:sheets/main:sheet", ns):
      sheet_name = sheet_info.attrib.get("name", "")
      if sheet_name.strip().lower() == "summary":
        ignored_sheets.append(sheet_name)
        continue
      target = relationships.get(sheet_info.attrib.get(f"{{{ns['rel']}}}id"))
      if not target:
        ignored_sheets.append(sheet_name)
        continue
      sheet_path = target.lstrip("/") if target.startswith("/") else "xl/" + target
      sheet_path = re.sub(r"^xl/\.\./", "", sheet_path)
      if sheet_path not in names:
        ignored_sheets.append(sheet_name)
        continue

      sheet_root = ET.fromstring(archive.read(sheet_path))
      grid = {}
      for row_node in sheet_root.findall(".//main:sheetData/main:row", ns):
        row_num = int(row_node.attrib.get("r", "0"))
        row_values = {}
        for cell in row_node.findall("main:c", ns):
          ref = cell.attrib.get("r", "")
          match = re.match(r"([A-Z]+)", ref)
          if not match:
            continue
          col_num = 0
          for char in match.group(1):
            col_num = col_num * 26 + ord(char) - 64
          cell_type = cell.attrib.get("t")
          value_node = cell.find("main:v", ns)
          if cell_type == "inlineStr":
            value = "".join(node.text or "" for node in cell.findall(".//main:t", ns))
          elif value_node is None:
            value = None
          elif cell_type == "s":
            try:
              value = shared_strings[int(value_node.text)]
            except (ValueError, IndexError, TypeError):
              value = None
          else:
            value = value_node.text
          row_values[col_num] = value
        if row_values:
          grid[row_num] = row_values

      header_num = None
      header = {}
      for row_num, row_values in grid.items():
        normalized = {
            col: str(value).strip().lower().replace("\n", " ") if value is not None else ""
            for col, value in row_values.items()
        }
        if "date" in normalized.values() and "hours" in normalized.values():
          header_num = row_num
          header = {value: col for col, value in normalized.items()}
          break
      if header_num is None:
        ignored_sheets.append(sheet_name)
        continue

      employee_name = sheet_name
      for row_values in grid.values():
        if str(row_values.get(1, "")).strip().lower() == "name:" and row_values.get(3):
          employee_name = str(row_values[3]).strip()
          break

      def get_cell(row_values, label):
        return row_values.get(header.get(label)) if header.get(label) else None

      def as_excel_datetime(value):
        try:
          return datetime.datetime(1899, 12, 30) + datetime.timedelta(days=float(value))
        except (TypeError, ValueError, OverflowError):
          return pd.to_datetime(value, errors="coerce")

      def as_excel_time(value):
        try:
          seconds = round(float(value) * 86400) % 86400
          return datetime.time(seconds // 3600, (seconds % 3600) // 60, seconds % 60)
        except (TypeError, ValueError, OverflowError):
          return value

      for row_num, row_values in grid.items():
        if row_num <= header_num:
          continue
        raw_date = get_cell(row_values, "date")
        if raw_date is None:
          continue
        parsed_date = as_excel_datetime(raw_date)
        if pd.isna(parsed_date):
          continue
        try:
          hours = float(get_cell(row_values, "hours") or 0)
        except (TypeError, ValueError):
          hours = 0.0
        try:
          advance = float(get_cell(row_values, "salary advance") or 0)
        except (TypeError, ValueError):
          advance = 0.0
        records.append({
            "الموظف": employee_name,
            "التاريخ": pd.Timestamp(parsed_date).normalize(),
            "اليوم": as_excel_datetime(get_cell(row_values, "day")) if get_cell(row_values, "day") is not None else None,
            "الحضور": as_excel_time(get_cell(row_values, "start time")),
            "الانصراف": as_excel_time(get_cell(row_values, "end time")),
            "ساعات العمل": hours,
            "السلفة (ج.م)": advance,
            "الفريق": "OneSpot",
        })

  return records, ignored_sheets


df_all, data_error = load_data()
data_blocked = data_error is not None

if data_error:
  st.error(data_error)
if not read_database_url():
  st.info(
      "Neon غير متصل بعد. التطبيق يعمل حاليًا بملف CSV محلي؛ أضف DATABASE_URL في Secrets "
      "على Streamlit لتخزين القيود في قاعدة البيانات."
  )


# =====================================================
# 4. واجهة OneSpot Cafe المالية
# =====================================================
st.markdown(
    """
    <style>
    html, body, [data-testid="stAppViewContainer"] {
        font-family: Arial, "Segoe UI", sans-serif !important;
    }
    [data-testid="stAppViewContainer"] p,
    [data-testid="stAppViewContainer"] h1,
    [data-testid="stAppViewContainer"] h2,
    [data-testid="stAppViewContainer"] h3,
    [data-testid="stAppViewContainer"] label,
    [data-testid="stAppViewContainer"] button,
    [data-testid="stAppViewContainer"] input,
    [data-testid="stAppViewContainer"] textarea,
    [data-testid="stAppViewContainer"] select,
    [data-testid="stAppViewContainer"] [data-testid="stMarkdownContainer"] {
        font-family: Arial, "Segoe UI", sans-serif !important;
    }
    .stApp { background: #f3f0eb; color: #2d2924; }
    [data-testid="stHeader"] { background: rgba(0, 0, 0, 0); }
    [data-testid="stMain"] { background: #f3f0eb; }
    [data-testid="stMainBlockContainer"] {
        max-width: 1480px;
        padding: 1.4rem 2rem 3rem;
    }
    [data-testid="stAppViewContainer"] { background: #f3f0eb; }
    .onespot-nav {
        background: #3b281b; color: #fff; border-radius: 8px 8px 0 0;
        padding: 13px 22px; display:flex; align-items:center; justify-content:flex-start;
        direction:rtl; text-align:right; margin: -1.4rem -2rem 0; min-height: 62px;
    }
    .onespot-brand { font-size: 19px; font-weight: 700; color:#fff; }
    .onespot-sub { font-size: 11px; letter-spacing:1.2px; color:#e8d2b8; margin-top:4px; }
    .onespot-links { color:#f3e9dd; font-size:13px; word-spacing:18px; }
    [data-baseweb="tab-list"] {
        direction:rtl; justify-content:flex-start; gap: 8px;
        background:#3b281b; border-bottom:0; margin: 0 -2rem 1.5rem;
        padding: 0 1.5rem 8px; border-radius:0 0 8px 8px;
    }
    [data-baseweb="tab"] {
        height: 47px; padding: 0 20px; color: #f3e9dd;
        background: transparent; border-radius: 8px;
        font-size: 16px !important; font-weight: 700 !important;
    }
    [data-baseweb="tab"][aria-selected="true"] {
        color: #fff; background: #584333; border-bottom: 2px solid #bc8848;
        font-size: 17px !important; font-weight: 800 !important;
    }
    [data-testid="stTabs"] [role="tab"],
    [data-testid="stTabs"] [role="tab"] p,
    [data-testid="stTabs"] [role="tab"] div {
        font-size: 18px !important;
        font-weight: 700 !important;
        line-height: 1.35 !important;
    }
    [data-testid="stTabs"] [role="tab"][aria-selected="true"],
    [data-testid="stTabs"] [role="tab"][aria-selected="true"] p {
        font-size: 19px !important;
        font-weight: 800 !important;
    }
    h1, h2, h3 { color:#2d2119; }
    [data-testid="stMetric"] {
        background: #ffffff; border: 1px solid #e7ddd2;
        border-radius: 12px; padding: 16px 18px; box-shadow: 0 2px 8px #3b281b0a;
    }
    [data-testid="stVerticalBlockBorderWrapper"] {
        border-color: #e7ddd2 !important; border-radius: 12px !important; background:#fff;
    }
    [data-testid="stRadio"] [role="radiogroup"] { gap: 8px; }
    [data-testid="stRadio"] [role="radiogroup"] label {
        background:#fff; border:1px solid #e5d9cc; border-radius:9px; padding:8px 14px;
    }
    [data-testid="stRadio"] [role="radiogroup"] label:has(input:checked) {
        background:#f1e8dc; border-color:#94652f;
    }
    [data-testid="stWidgetLabel"] p { color:#6f5a48; }
    [data-testid="stTextInput"] input, [data-testid="stNumberInput"] input,
    [data-testid="stDateInput"] input, [data-testid="stSelectbox"] [data-baseweb="select"] > div {
        border-radius:8px !important; border-color:#e5d9cc !important; background:#fbf9f6;
    }
    [data-testid="stFormSubmitButton"] button,
    [data-testid="stButton"] button[kind="primary"], button[kind="primary"] {
        background:#bc8848 !important; color:#fff !important; border:1px solid #bc8848 !important;
        border-radius:8px !important; font-weight:700 !important;
    }
    [data-testid="stFormSubmitButton"] button:hover { background:#94652f !important; border-color:#94652f !important; }
    [data-testid="stDataFrame"] { border:1px solid #e7ddd2; border-radius:10px; overflow:hidden; }
    hr { border-color:#ddd2c7 !important; }
    </style>
    """,
    unsafe_allow_html=True,
)

st.markdown(
    '<div class="onespot-nav"><div><div class="onespot-brand">☕ OneSpot Cafe</div>'
    '<div class="onespot-sub">COFFEE &amp; MORE · FINANCIAL SYSTEM</div></div></div>',
    unsafe_allow_html=True,
)
tab_home, tab_add, tab_history, tab_admin, tab_reports, tab_attendance = st.tabs(
    ["🏠 الرئيسية", "➕ إضافة قيد", "📒 سجل المعاملات", "🗂️ إدارة القيود", "📊 مقارنة وتقارير", "🕒 الحضور والانصراف"]
)

# =====================================================
# 5. تسجيل قيد جديد
# =====================================================
if "form_version" not in st.session_state:
    st.session_state["form_version"] = 0


@st.fragment
def transaction_form_fragment():
    st.subheader("＋ إضافة قيد جديد")

    if data_blocked:
        st.warning("تم إيقاف تسجيل القيود حتى إصلاح ملف البيانات.")
        return

    version = st.session_state["form_version"]

    with st.container(border=True):
        tx_type = st.radio(
            "نوع القيد:",
            ["Income", "Expense"],
            index=None,
            horizontal=True,
            format_func=lambda value: "إيراد" if value == "Income" else "مصروف",
            key=f"tp_{version}",
        )

        if tx_type is None:
            st.info("اختر إيراد أو مصروف، ثم أدخل بيانات القيد.")

        with st.form(key=f"transaction_form_{version}"):
            branch_col, date_col, category_col, amount_col, description_col = st.columns([.9, 1.05, 1.15, .9, 1.9])
            with branch_col:
                transaction_branch = st.selectbox(
                    "الفرع:",
                    [None] + BRANCHES,
                    index=0,
                    format_func=lambda value: "اختر الفرع" if value is None else value,
                    key=f"branch_{version}",
                )
            with date_col:
                tx_date = st.date_input(
                    "التاريخ:",
                    value=datetime.date.today(),
                    key=f"dt_{version}",
                )
            with category_col:
                if tx_type == "Income":
                    tx_category = st.selectbox(
                        "البند المحاسبي:",
                        INCOME_CATS,
                        key=f"cat_in_{version}",
                    )
                elif tx_type == "Expense":
                    tx_category = st.selectbox(
                        "البند المحاسبي:",
                        EXPENSE_CATS,
                        key=f"cat_ex_{version}",
                    )
                else:
                    tx_category = None
                    st.selectbox(
                        "البند المحاسبي:",
                        ["اختر نوع القيد أولاً"],
                        disabled=True,
                        key=f"cat_none_{version}",
                    )
            with amount_col:
                tx_amount = st.number_input(
                    "المبلغ (ج.م):",
                    min_value=0.0,
                    step=10.0,
                    format="%.2f",
                    key=f"amt_{version}",
                )
            with description_col:
                tx_description = st.text_input(
                    "البيان:",
                    key=f"desc_{version}",
                )
            st.caption("اضغط Enter من خانة البيان للحفظ السريع.")

            submitted = st.form_submit_button(
                "حفظ القيد",
                type="primary",
                use_container_width=True,
            )

    if not submitted:
        return

    if tx_type not in ["Income", "Expense"]:
        st.error("اختر نوع المعاملة.")
    elif transaction_branch not in BRANCHES:
        st.error("اختر الفرع قبل الحفظ.")
    elif tx_category is None:
        st.error("اختر البند المحاسبي.")
    elif tx_amount <= 0:
        st.error("يجب أن يكون المبلغ أكبر من صفر.")
    elif not tx_description.strip():
        st.error("اكتب البيان قبل الحفظ.")
    else:
        latest_df, latest_error = load_data()
        if latest_error:
            st.error("تعذر التحقق من ملف البيانات: " + latest_error)
            return

        next_id = int(latest_df["ID"].max()) + 1 if not latest_df.empty else 1
        new_row = pd.DataFrame([{
            "ID": next_id,
            "Date": tx_date.strftime("%Y-%m-%d"),
            "Month_No": tx_date.month,
            "Branch": transaction_branch,
            "Category": tx_category,
            "Type": tx_type,
            "Amount": float(tx_amount),
            "Description": tx_description.strip(),
        }])
        updated_df = pd.concat([latest_df, new_row], ignore_index=True)

        try:
            save_data(updated_df)
            st.session_state["form_version"] += 1
            st.success(f"تم حفظ القيد رقم {next_id} بنجاح.")
            st.rerun()
        except Exception as exc:
            st.error(f"تعذر حفظ القيد والحفاظ على البيانات. التفاصيل: {exc}")


with tab_add:
    transaction_form_fragment()

st.markdown("---")

# Use an empty, schema-compatible frame when the CSV is unavailable.
data_view = df_all if df_all is not None else pd.DataFrame(columns=COLUMNS)

with tab_home:
    # =====================================================
    # 6. الملخص المالي والرسوم البيانية
    # =====================================================
    st.subheader("ملخص الأداء المالي")
    
    if not data_view.empty:
        report_df = data_view.copy()
        report_df["Amount"] = pd.to_numeric(report_df["Amount"], errors="coerce").fillna(0)
        report_df["Parsed_Date"] = pd.to_datetime(report_df["Date"], errors="coerce")
        report_df["Year_Month"] = report_df["Parsed_Date"].dt.strftime("%Y-%m")
    
        selected_branch = st.radio(
            "عرض بيانات:",
            ["إجمالي الفرعين"] + BRANCHES,
            horizontal=True,
            key="dashboard_branch_filter",
        )
        analysis_df = (
            report_df if selected_branch == "إجمالي الفرعين"
            else report_df[report_df["Branch"] == selected_branch]
        )
        total_income = analysis_df.loc[analysis_df["Type"] == "Income", "Amount"].sum()
        total_expense = analysis_df.loc[analysis_df["Type"] == "Expense", "Amount"].sum()
        net_profit = total_income - total_expense
        summary_scope = "الفرعين" if selected_branch == "إجمالي الفرعين" else selected_branch
    
        metric_income, metric_expense, metric_profit = st.columns(3)
        metric_income.metric(f"إيرادات {summary_scope}", format_currency(total_income))
        metric_expense.metric(f"مصروفات {summary_scope}", format_currency(total_expense))
        metric_profit.metric(
            f"صافي ربح {summary_scope}",
            format_currency(net_profit),
            delta=format_currency(net_profit),
        )

        st.markdown("### ملخص الأداء المالي لكل فرع")
        branch_columns = st.columns(len(BRANCHES))
        for branch_column, branch_name in zip(branch_columns, BRANCHES):
            branch_data = report_df[report_df["Branch"] == branch_name]
            branch_income = branch_data.loc[branch_data["Type"] == "Income", "Amount"].sum()
            branch_expense = branch_data.loc[branch_data["Type"] == "Expense", "Amount"].sum()
            with branch_column:
                with st.container(border=True):
                    st.markdown(f"#### {branch_name}")
                    st.metric("الإيرادات", format_currency(branch_income))
                    st.metric("المصروفات", format_currency(branch_expense))
                    st.metric("صافي الربح", format_currency(branch_income - branch_expense))
                    st.caption(f"عدد القيود المسجلة: {len(branch_data)}")

        chart_col1, chart_col2 = st.columns(2)
        with chart_col1:
            with st.container(border=True):
                st.markdown("**توزيع المصروفات حسب البند**")
                expense_df = analysis_df[analysis_df["Type"] == "Expense"]
                if HAS_PLOTLY and not expense_df.empty:
                    fig_pie = px.pie(
                        expense_df,
                        names="Category",
                        values="Amount",
                        hole=0.48,
                        color_discrete_sequence=["#bc8848", "#557b60", "#b85045", "#847263"],
                    )
                    fig_pie.update_traces(
                        textposition="inside",
                        textinfo="percent",
                        hovertemplate="%{label}<br>%{value:,.2f} ج.م (%{percent})<extra></extra>",
                    )
                    fig_pie.update_layout(
                        height=285,
                        margin=dict(l=8, r=8, t=12, b=8),
                        showlegend=True,
                        legend=dict(orientation="h", y=-0.12),
                        paper_bgcolor="rgba(0,0,0,0)",
                        plot_bgcolor="rgba(0,0,0,0)",
                    )
                    st.plotly_chart(fig_pie, use_container_width=True)
                elif expense_df.empty:
                    st.info("ستظهر المصروفات هنا بعد تسجيل أول مصروف.")
                else:
                    st.info("ثبّت Plotly لعرض الرسم البياني.")
    
        with chart_col2:
            with st.container(border=True):
                st.markdown("**الإيرادات والمصروفات شهرياً**")
                monthly_trend = (
                    analysis_df.groupby(["Year_Month", "Type"])["Amount"]
                    .sum()
                    .reset_index()
                )
                if HAS_PLOTLY and not monthly_trend.empty:
                    fig_bar = px.bar(
                        monthly_trend,
                        x="Year_Month",
                        y="Amount",
                        color="Type",
                        barmode="group",
                        color_discrete_map={"Income": "#557b60", "Expense": "#b85045"},
                        labels={"Year_Month": "الشهر", "Amount": "المبلغ (ج.م)", "Type": "النوع"},
                    )
                    fig_bar.update_traces(hovertemplate="%{y:,.2f} ج.م<extra></extra>")
                    fig_bar.update_layout(
                        height=285,
                        margin=dict(l=8, r=8, t=12, b=8),
                        legend_title_text="",
                        paper_bgcolor="rgba(0,0,0,0)",
                        plot_bgcolor="rgba(0,0,0,0)",
                    )
                    fig_bar.update_yaxes(title_text="المبلغ (ج.م)", gridcolor="#e7e1d9")
                    st.plotly_chart(fig_bar, use_container_width=True)
                elif not HAS_PLOTLY:
                    st.info("ثبّت Plotly لعرض الرسم البياني.")
                else:
                    st.info("ستظهر المقارنة بعد تسجيل المعاملات.")
    
        pivot_df = analysis_df.pivot_table(
            index="Year_Month",
            columns="Type",
            values="Amount",
            aggfunc="sum",
            fill_value=0,
        )
        for col in ["Income", "Expense"]:
            if col not in pivot_df.columns:
                pivot_df[col] = 0.0
    
        pivot_df["Net Profit"] = pivot_df["Income"] - pivot_df["Expense"]
        pivot_df = pivot_df.sort_index(ascending=False)
        summary = pivot_df.rename(
            columns={"Income": "إجمالي الإيرادات", "Expense": "إجمالي المصروفات"}
        )[["إجمالي الإيرادات", "إجمالي المصروفات", "Net Profit"]]
        summary = summary.rename_axis("الشهر").rename(columns={"Net Profit": "صافي الربح"})
        st.dataframe(
            summary.style.format({
                "إجمالي الإيرادات": "{:,.2f} ج.م",
                "إجمالي المصروفات": "{:,.2f} ج.م",
                "صافي الربح": "{:,.2f} ج.م",
            }),
            use_container_width=True,
        )
    else:
        st.info("سجّل أول قيد لعرض ملخص الإيرادات والمصروفات والرسوم.")
    
    st.markdown("---")
    
with tab_history:
    # =====================================================
    # 7. سجل المعاملات والبحث
    # =====================================================
    st.subheader("سجل المعاملات")
    
    if data_view.empty:
        st.info("لا توجد قيود مسجلة حتى الآن.")
    else:
        with st.expander("تصفية وبحث في القيود", expanded=False):
            filter_col1, filter_col2 = st.columns(2)
            with filter_col1:
                min_date = pd.to_datetime(data_view["Date"]).min().date()
                max_date = pd.to_datetime(data_view["Date"]).max().date()
                start_date = st.date_input(
                    "من تاريخ:",
                    value=min_date,
                    key="start_dt_tab1",
                )
            with filter_col2:
                end_date = st.date_input(
                    "إلى تاريخ:",
                    value=max_date,
                    key="end_dt_tab1",
                )
    
            filter_col3, filter_col4, filter_col5, filter_col6 = st.columns(4)
            with filter_col3:
                filter_type = st.selectbox(
                    "نوع المعاملة:",
                    ["All", "Income", "Expense"],
                    format_func=lambda value: {
                        "All": "الكل",
                        "Income": "إيراد",
                        "Expense": "مصروف",
                    }[value],
                    key="flt_type_t1",
                )
            with filter_col4:
                all_categories = ["All"] + list(data_view["Category"].unique())
                filter_category = st.selectbox(
                    "البند:",
                    all_categories,
                    format_func=lambda value: "الكل" if value == "All" else value,
                    key="flt_cat_t1",
                )
            with filter_col5:
                search_description = st.text_input(
                    "بحث في البيان:",
                    key="search_desc_t1",
                )
            with filter_col6:
                branch_filter = st.selectbox(
                    "الفرع:",
                    ["All"] + BRANCHES,
                    format_func=lambda value: "الكل" if value == "All" else value,
                    key="flt_branch_t1",
                )
    
        filtered_df = data_view.copy()
        filtered_df["Parsed_Date_Obj"] = pd.to_datetime(
            filtered_df["Date"], errors="coerce"
        ).dt.date
        filtered_df = filtered_df[
            (filtered_df["Parsed_Date_Obj"] >= start_date)
            & (filtered_df["Parsed_Date_Obj"] <= end_date)
        ]
        if filter_type != "All":
            filtered_df = filtered_df[filtered_df["Type"] == filter_type]
        if filter_category != "All":
            filtered_df = filtered_df[filtered_df["Category"] == filter_category]
        if branch_filter != "All":
            filtered_df = filtered_df[filtered_df["Branch"] == branch_filter]
        if search_description.strip():
            filtered_df = filtered_df[
                filtered_df["Description"].str.contains(
                    search_description,
                    case=False,
                    na=False,
                )
            ]
    
        csv_data = (
            filtered_df.drop(columns=["Parsed_Date_Obj"])
            .to_csv(index=False)
            .encode("utf-8-sig")
        )
        st.download_button(
            label="تنزيل القيود المعروضة بصيغة CSV",
            data=csv_data,
            file_name=f"onespot_transactions_{start_date}_to_{end_date}.csv",
            mime="text/csv",
        )
    
        def show_branch_transactions(branch_frame, branch_name):
            st.markdown(f"#### {branch_name}")
            if branch_frame.empty:
                st.info(f"لا توجد قيود مطابقة لفرع {branch_name}.")
                return

            display_df = branch_frame.sort_values("ID", ascending=False).copy()
            display_df["Amount"] = display_df["Amount"].map(format_currency)
            display_df = display_df.rename(columns={
                "ID": "رقم القيد",
                "Date": "التاريخ",
                "Branch": "الفرع",
                "Category": "البند",
                "Type": "النوع",
                "Amount": "المبلغ",
                "Description": "البيان",
            })
            display_df["النوع"] = display_df["النوع"].map({
                "Income": "إيراد",
                "Expense": "مصروف",
            })
            display_df = display_df.drop(
                columns=["Month_No", "Parsed_Date_Obj"], errors="ignore"
            )
            st.dataframe(
                display_df.head(10),
                use_container_width=True,
                hide_index=True,
            )
            if len(display_df) > 10:
                st.caption("يعرض الجدول أحدث 10 قيود لهذا الفرع. استخدم البحث أو تنزيل CSV لبقية النتائج.")

        branches_to_show = BRANCHES if branch_filter == "All" else [branch_filter]
        for branch_name in branches_to_show:
            branch_frame = filtered_df[filtered_df["Branch"] == branch_name]
            show_branch_transactions(branch_frame, branch_name)
    
with tab_admin:
    # =====================================================
    # 8. حذف القيود
    # =====================================================
    with st.expander("مسح القيد", expanded=True):
        if data_blocked:
            st.error("الحذف متوقف حتى إصلاح مشكلة ملف البيانات.")
        elif data_view.empty:
            st.info("لا توجد قيود متاحة للحذف.")
        else:
            delete_df = data_view.sort_values("ID", ascending=False).copy()
            delete_options = {
                int(row["ID"]): (
                    f"رقم {int(row['ID'])} | {row['Date']} | "
                f"{row['Branch']} | "
                    f"{'إيراد' if row['Type'] == 'Income' else 'مصروف'} | "
                    f"{format_currency(row['Amount'])} | "
                    f"{str(row['Description'])[:40]}"
                )
                for _, row in delete_df.iterrows()
            }
            selected_id = st.selectbox(
                "اختر القيد الذي تريد حذفه:",
                options=[None] + list(delete_options.keys()),
                format_func=lambda value: (
                    "-- اختر القيد --" if value is None else delete_options[value]
                ),
                key="tab_del_select",
            )
    
            if selected_id is not None:
                selected_rows = data_view.loc[data_view["ID"] == selected_id]
                if not selected_rows.empty:
                    row = selected_rows.iloc[0]
                    st.warning(
                        f"القيد المحدد: رقم {int(row['ID'])} | {row['Date']} | "
                    f"{row['Branch']} | "
                        f"{'إيراد' if row['Type'] == 'Income' else 'مصروف'} | "
                        f"{format_currency(row['Amount'])} | {row['Description']}"
                    )
                    confirm_delete = st.checkbox(
                        "أؤكد أنني أريد حذف هذا القيد نهائيًا.",
                        key=f"confirm_delete_{selected_id}",
                    )
                    if st.button(
                        "تأكيد حذف القيد",
                        type="primary",
                        disabled=not confirm_delete,
                        key=f"delete_btn_{selected_id}",
                    ):
                        latest_df, latest_error = load_data()
                        if latest_error:
                            st.error("تم إيقاف الحذف لحماية البيانات: " + latest_error)
                        elif selected_id not in latest_df["ID"].values:
                            st.error("القيد غير موجود في الملف الحالي. أعد تحميل الصفحة.")
                        else:
                            updated_df = latest_df.loc[
                                latest_df["ID"] != selected_id
                            ].copy()
                            try:
                                save_data(updated_df)
                                st.success(
                                    f"تم حذف القيد رقم {selected_id} وتحديث البيانات."
                                )
                                st.rerun()
                            except Exception as exc:
                                st.error(f"تعذر إكمال الحذف والحفظ. التفاصيل: {exc}")

with tab_reports:
    st.subheader("مقارنة الفروع والتقرير الشهري")
    if data_view.empty:
        st.info("سجّل قيودًا لتظهر مقارنة الفروع والتقارير هنا.")
    else:
        report_view = data_view.copy()
        report_view["Amount"] = pd.to_numeric(report_view["Amount"], errors="coerce").fillna(0)
        report_view["Parsed_Date"] = pd.to_datetime(report_view["Date"], errors="coerce")
        report_view["Year_Month"] = report_view["Parsed_Date"].dt.strftime("%Y-%m")
        report_months = sorted(report_view["Year_Month"].dropna().unique(), reverse=True)
        chosen_month = st.selectbox(
            "اختاري الشهر للتقرير:",
            ["كل الشهور"] + report_months,
            key="branch_report_month",
        )
        if chosen_month != "كل الشهور":
            report_view = report_view[report_view["Year_Month"] == chosen_month].copy()

        comparison_rows = []
        for branch_name in BRANCHES:
            branch_data = report_view[report_view["Branch"] == branch_name]
            branch_income = branch_data.loc[branch_data["Type"] == "Income", "Amount"].sum()
            branch_expense = branch_data.loc[branch_data["Type"] == "Expense", "Amount"].sum()
            comparison_rows.append({
                "الفرع": branch_name,
                "الإيرادات": branch_income,
                "المصروفات": branch_expense,
                "صافي الربح": branch_income - branch_expense,
                "عدد القيود": len(branch_data),
            })

        comparison_total = {
            "الفرع": "إجمالي الفرعين",
            "الإيرادات": sum(row["الإيرادات"] for row in comparison_rows),
            "المصروفات": sum(row["المصروفات"] for row in comparison_rows),
            "صافي الربح": sum(row["صافي الربح"] for row in comparison_rows),
            "عدد القيود": sum(row["عدد القيود"] for row in comparison_rows),
        }
        comparison_rows.append(comparison_total)
        st.dataframe(
            pd.DataFrame(comparison_rows).style.format({
                "الإيرادات": "{:,.2f} ج.م",
                "المصروفات": "{:,.2f} ج.م",
                "صافي الربح": "{:,.2f} ج.م",
                "عدد القيود": "{:,.0f}",
            }),
            use_container_width=True,
            hide_index=True,
        )

        st.markdown("#### تنبيه عند تجاوز مصروفات الفرع")
        alert_columns = st.columns(len(BRANCHES))
        for alert_column, branch_name in zip(alert_columns, BRANCHES):
            branch_expenses = next(
                row["المصروفات"] for row in comparison_rows if row["الفرع"] == branch_name
            )
            with alert_column:
                expense_limit = st.number_input(
                    f"حد المصروفات لـ {branch_name} (ج.م):",
                    min_value=0.0,
                    value=0.0,
                    step=500.0,
                    key=f"expense_alert_{branch_name}",
                )
                if expense_limit > 0 and branch_expenses > expense_limit:
                    st.warning(
                        f"مصروفات {branch_name} تجاوزت الحد المحدد لهذا الشهر."
                    )

        if HAS_PLOTLY and not report_view.empty:
            comparison_chart_data = (
                report_view.groupby(["Branch", "Type"])["Amount"].sum().reset_index()
            )
            comparison_chart = px.bar(
                comparison_chart_data,
                x="Branch",
                y="Amount",
                color="Type",
                barmode="group",
                color_discrete_map={"Income": "#557b60", "Expense": "#b85045"},
                labels={"Branch": "الفرع", "Amount": "المبلغ (ج.م)", "Type": "النوع"},
            )
            comparison_chart.update_layout(
                height=360,
                margin=dict(l=12, r=12, t=20, b=12),
                legend_title_text="",
                paper_bgcolor="rgba(0,0,0,0)",
                plot_bgcolor="rgba(0,0,0,0)",
            )
            comparison_chart.update_traces(hovertemplate="%{y:,.2f} ج.م<extra></extra>")
            st.plotly_chart(comparison_chart, use_container_width=True)

        export_df = report_view.drop(columns=["Month_No", "Parsed_Date", "Year_Month"], errors="ignore")
        export_df = export_df.rename(columns={
            "ID": "رقم القيد",
            "Date": "التاريخ",
            "Branch": "الفرع",
            "Category": "البند",
            "Type": "النوع",
            "Amount": "المبلغ بالجنيه",
            "Description": "البيان",
        })
        export_df["النوع"] = export_df["النوع"].map({"Income": "إيراد", "Expense": "مصروف"})
        st.download_button(
            "تنزيل التقرير بصيغة CSV",
            data=export_df.to_csv(index=False).encode("utf-8-sig"),
            file_name=f"onespot_branch_report_{chosen_month.replace(' ', '_')}.csv",
            mime="text/csv",
        )

with tab_attendance:
    st.subheader("الحضور والانصراف — فريق OneSpot")
    st.caption(
        "ارفع ملف الحضور بصيغة Excel لعرض ساعات العمل والسلف لكل موظف. "
        "سيتم جمع الموظفين معًا تحت OneSpot لأن مواقع عملهم متغيرة."
    )

    attendance_file = st.file_uploader(
        "اختار ملف الحضور والانصراف (.xlsm أو .xlsx)",
        type=["xlsm", "xlsx"],
        key="attendance_workbook_upload",
        help="يتم قراءة أوراق الموظفين فقط، وتجاهل ورقة Summary حتى لا تتكرر البيانات.",
    )

    if attendance_file is not None:
        try:
            from openpyxl import load_workbook

            attendance_book = load_workbook(
                attendance_file, read_only=True, data_only=True, keep_vba=False
            )
            attendance_rows = []
            ignored_sheets = []
            for sheet in attendance_book.worksheets:
                if sheet.title.strip().lower() == "summary":
                    ignored_sheets.append(sheet.title)
                    continue
                header_row = None
                headers = []
                for row_num, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    normalized = [str(value).strip().lower().replace("\n", " ") if value is not None else "" for value in row]
                    if "date" in normalized and "hours" in normalized:
                        header_row = row_num
                        headers = normalized
                        break

                if header_row is None:
                    ignored_sheets.append(sheet.title)
                    continue

                positions = {name: headers.index(name) for name in ["date", "day", "salary advance", "start time", "end time", "hours"] if name in headers}
                employee_name = None
                for row in sheet.iter_rows(min_row=1, max_row=min(header_row, 8), values_only=True):
                    if len(row) > 2 and str(row[0]).strip().lower() == "name:" and row[2]:
                        employee_name = str(row[2]).strip()
                        break
                employee_name = employee_name or sheet.title

                for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
                    date_value = row[positions["date"]] if positions["date"] < len(row) else None
                    if date_value is None:
                        continue
                    date_value = pd.to_datetime(date_value, errors="coerce")
                    if pd.isna(date_value):
                        continue
                    def cell_value(field):
                        idx = positions.get(field)
                        return row[idx] if idx is not None and idx < len(row) else None

                    hours = pd.to_numeric(cell_value("hours"), errors="coerce")
                    advance = pd.to_numeric(cell_value("salary advance"), errors="coerce")
                    attendance_rows.append({
                        "الموظف": employee_name,
                        "التاريخ": date_value.normalize(),
                        "اليوم": cell_value("day"),
                        "الحضور": cell_value("start time"),
                        "الانصراف": cell_value("end time"),
                        "ساعات العمل": 0.0 if pd.isna(hours) else float(hours),
                        "السلفة (ج.م)": 0.0 if pd.isna(advance) else float(advance),
                        "الفريق": "OneSpot",
                    })

            attendance_book.close()
            if not attendance_rows:
                st.warning("لم أجد سجلات حضور. تأكد أن أوراق الموظفين تحتوي على أعمدة Date وHours.")
            else:
                attendance_df = pd.DataFrame(attendance_rows)
                attendance_df = attendance_df.drop_duplicates(
                    subset=["الموظف", "التاريخ"], keep="last"
                ).sort_values(["التاريخ", "الموظف"], ascending=[False, True])
                save_attendance_data(attendance_df)
                st.session_state["attendance_imported_df"] = attendance_df
                if ignored_sheets:
                    st.caption("تم تجاهل الأوراق غير الخاصة بسجلات الموظفين: " + "، ".join(ignored_sheets))
        except ImportError:
            try:
                attendance_rows, ignored_sheets = read_attendance_with_builtin_excel(attendance_file)
                if attendance_rows:
                    attendance_df = pd.DataFrame(attendance_rows)
                    attendance_df = attendance_df.drop_duplicates(
                        subset=["الموظف", "التاريخ"], keep="last"
                    ).sort_values(["التاريخ", "الموظف"], ascending=[False, True])
                    save_attendance_data(attendance_df)
                    st.session_state["attendance_imported_df"] = attendance_df
                    if ignored_sheets:
                        st.caption("تم تجاهل الأوراق غير الخاصة بسجلات الموظفين: " + "، ".join(ignored_sheets))
                else:
                    st.warning("لم أجد سجلات حضور في أوراق الموظفين.")
            except Exception as exc:
                st.error(f"تعذر قراءة ملف الحضور. تأكد أن الملف سليم ثم ارفعه مرة أخرى. التفاصيل: {exc}")
        except Exception as exc:
            st.error(f"تعذر قراءة ملف الحضور. تأكد أن الملف سليم ثم ارفعه مرة أخرى. التفاصيل: {exc}")

    attendance_df = st.session_state.get("attendance_imported_df")
    if read_database_url():
        try:
            persisted_attendance = load_attendance_data()
            if not persisted_attendance.empty:
                attendance_df = persisted_attendance
                st.session_state["attendance_imported_df"] = persisted_attendance
        except Exception as exc:
            st.warning(f"تعذر تحميل بيانات الحضور المحفوظة من Neon: {exc}")
    if attendance_df is not None and not attendance_df.empty:
        month_values = sorted(attendance_df["التاريخ"].dt.strftime("%Y-%m").unique(), reverse=True)
        selected_attendance_month = st.selectbox(
            "الشهر:", ["كل الشهور"] + month_values, key="attendance_month_filter"
        )
        attendance_view = attendance_df.copy()
        if selected_attendance_month != "كل الشهور":
            attendance_view = attendance_view[
                attendance_view["التاريخ"].dt.strftime("%Y-%m") == selected_attendance_month
            ]

        total_hours = attendance_view["ساعات العمل"].sum()
        total_advances = attendance_view["السلفة (ج.م)"].sum()
        employees_count = attendance_view["الموظف"].nunique()
        worked_days = int((attendance_view["ساعات العمل"] > 0).sum())
        metric_cols = st.columns(4)
        metric_cols[0].metric("عدد الموظفين", f"{employees_count}")
        metric_cols[1].metric("إجمالي ساعات العمل", f"{total_hours:,.2f} ساعة")
        metric_cols[2].metric("أيام العمل المسجلة", f"{worked_days}")
        metric_cols[3].metric("إجمالي السلف", format_currency(total_advances))

        employee_summary = attendance_view.groupby("الموظف", as_index=False).agg(
            **{"أيام العمل": ("ساعات العمل", lambda values: int((values > 0).sum())),
               "إجمالي الساعات": ("ساعات العمل", "sum"),
               "إجمالي السلف (ج.م)": ("السلفة (ج.م)", "sum")}
        )
        st.markdown("#### ملخص الموظفين")
        st.dataframe(
            employee_summary.style.format({
                "أيام العمل": "{:,.0f}",
                "إجمالي الساعات": "{:,.2f}",
                "إجمالي السلف (ج.م)": "{:,.2f}",
            }),
            use_container_width=True,
            hide_index=True,
        )
        st.markdown("#### تفاصيل الحضور والانصراف")
        display_df = attendance_view.rename(columns={"التاريخ": "التاريخ"}).copy()
        display_df["التاريخ"] = display_df["التاريخ"].dt.strftime("%Y-%m-%d")
        st.dataframe(
            display_df[["الموظف", "التاريخ", "اليوم", "الحضور", "الانصراف", "ساعات العمل", "السلفة (ج.م)", "الفريق"]],
            use_container_width=True,
            hide_index=True,
        )
        st.download_button(
            "تنزيل بيانات الحضور كملف CSV",
            data=attendance_view.assign(**{"التاريخ": attendance_view["التاريخ"].dt.strftime("%Y-%m-%d")}).to_csv(index=False).encode("utf-8-sig"),
            file_name=f"onespot_attendance_{selected_attendance_month.replace(' ', '_')}.csv",
            mime="text/csv",
        )
    
